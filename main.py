"""Run HRS: read the booking spreadsheet, search CCIS for each person, save PDFs.

    python main.py <spreadsheet.xlsx>             search everyone
    python main.py <spreadsheet.xlsx> --limit 3   search the first 3 only
    python main.py <spreadsheet.xlsx> --check     read the sheet only, no CCIS

Console and log lines name spreadsheet rows, never people, so either is safe
to paste. The PDFs themselves are named LAST_FIRST.pdf, so the folder they
land in is gitignored.
"""

import argparse
import re
import sys
from collections import Counter
from datetime import date
from pathlib import Path

from openpyxl import load_workbook

from ccis import CcisSession
from config_loader import load_config
from excel_reader import HEADER_SCAN_ROWS, SUFFIXES, Person, find_columns, read_people
from exceptions import RowProblem, SystemProblem
from logger import get_logger, setup_logging

# Anything outside these is replaced in file names. Windows refuses some
# characters outright, and a stray one would fail the write for that person.
UNSAFE_FILENAME_CHARS = re.compile(r"[^A-Za-z0-9'-]+")

HISTORY = "criminal history"
NO_RESULTS = "no results found"
NEEDS_A_LOOK = "needs a look by hand"


def describe_layout(path: Path) -> None:
    """Print the header position, chosen columns and data row count."""
    workbook = load_workbook(path, read_only=True, data_only=True)
    try:
        sheet = workbook.worksheets[0]
        columns = find_columns(
            sheet.iter_rows(max_row=HEADER_SCAN_ROWS, values_only=True)
        )
        data_rows = sum(1 for _ in sheet.iter_rows(min_row=columns.header_row + 1))
        print(f"Sheets in workbook:  {len(workbook.worksheets)} (reading the first)")
        print(f"Header row:          {columns.header_row}")
        print(f"Name column:         {columns.name + 1}")
        print(f"DOB column:          {columns.dob + 1}")
        print(f"Rows below header:   {data_rows}")
    finally:
        workbook.close()


def name_shape(person: Person) -> str:
    """Describe a name's form without revealing it."""
    # A suffix after a second comma counts as part of the first-name side.
    last_part, first_part = person.raw_name.split(",", 1)
    first_part = first_part.replace(",", " ")
    tokens = [t.rstrip(".").upper() for t in (last_part + " " + first_part).split()]
    parts = [f"{len(last_part.split())}-word last", f"{len(first_part.split())}-word first"]
    if any(t in SUFFIXES for t in tokens):
        parts.append("suffix")
    return ", ".join(parts)


def describe_people(people: list[Person]) -> None:
    """Print how many people were found, the problem rows and name shapes."""
    problems = [p for p in people if p.problem]
    good = [p for p in people if not p.problem]

    print(f"People found:        {len(people)}")
    print(f"Ready to search:     {len(good)}")
    print(f"Flagged:             {len(problems)}")
    for person in problems:
        print(f"  row {person.row}: {person.problem}")

    print("Name shapes (count):")
    for shape, count in sorted(Counter(name_shape(p) for p in good).items()):
        print(f"  {count:>4}  {shape}")


def pdf_name(person: Person, used: set[str]) -> str:
    """Return LAST_FIRST.pdf, adding _2, _3 when two people in a run share it.

    A file left from an earlier run with the same name is overwritten, so a
    rerun replaces its own PDFs instead of piling up copies.
    """
    stem = "_".join(
        UNSAFE_FILENAME_CHARS.sub("-", part).strip("-").upper()
        for part in (person.last, person.first)
    )
    name = f"{stem}.pdf"
    copy = 2
    while name in used:
        name = f"{stem}_{copy}.pdf"
        copy += 1
    used.add(name)
    return name


def search_people(config: dict, people: list[Person]) -> dict[str, list[int]]:
    """Search CCIS for each person and save their PDF.

    Returns the spreadsheet rows under each outcome. A RowProblem sends that
    person to "needs a look by hand" and the run carries on; a SystemProblem
    stops it.
    """
    log = get_logger()
    pdf_dir: Path = config["paths"]["pdfs"]
    try:
        pdf_dir.mkdir(parents=True, exist_ok=True)
    except OSError as error:
        raise SystemProblem(f"Could not create the PDF folder: {error}") from error

    outcomes: dict[str, list[int]] = {HISTORY: [], NO_RESULTS: [], NEEDS_A_LOOK: []}
    used: set[str] = set()
    with CcisSession(config) as ccis:
        for number, person in enumerate(people, start=1):
            try:
                result = ccis.fetch_pdf(person.first, person.last, person.dob)
            except RowProblem as error:
                log.warning("Row %d (%d of %d): %s: %s",
                            person.row, number, len(people), NEEDS_A_LOOK, error)
                outcomes[NEEDS_A_LOOK].append(person.row)
                continue

            try:
                (pdf_dir / pdf_name(person, used)).write_bytes(result.pdf_bytes)
            except OSError as error:
                # A folder that cannot be written to fails every person after
                # this one too, so it stops the run.
                raise SystemProblem(
                    f"Could not save the PDF for row {person.row}: {type(error).__name__}"
                ) from error

            outcome = HISTORY if result.has_record else NO_RESULTS
            log.info("Row %d (%d of %d): %s", person.row, number, len(people), outcome)
            outcomes[outcome].append(person.row)
    return outcomes


def run(path: Path, limit: int | None) -> int:
    """Load config, read the sheet, search everyone ready, and print a summary."""
    config = load_config()
    setup_logging(config["paths"]["logs"] / f"{date.today():%Y-%m-%d}.log")
    log = get_logger()

    people = read_people(path)
    ready = [p for p in people if not p.problem]
    flagged = [p for p in people if p.problem]
    log.info("People found: %d, ready: %d, flagged: %d",
             len(people), len(ready), len(flagged))
    for person in flagged:
        log.warning("Row %d: not searched: %s", person.row, person.problem)
    if limit is not None:
        ready = ready[:limit]
        log.info("Limited to the first %d ready", len(ready))

    outcomes = search_people(config, ready)

    print()
    print(f"Searched:             {len(ready)}")
    for outcome, rows in outcomes.items():
        print(f"{outcome + ':':<22}{len(rows)}")
    for row in outcomes[NEEDS_A_LOOK]:
        print(f"  row {row}")
    print(f"Flagged, not searched: {len(flagged)}")
    for person in flagged:
        print(f"  row {person.row}: {person.problem}")
    return 0


def main() -> int:
    """Parse the command line and run, turning a SystemProblem into exit code 1."""
    parser = argparse.ArgumentParser(description="Search CCIS for everyone on a booking sheet.")
    parser.add_argument("spreadsheet", type=Path)
    parser.add_argument("--limit", type=int, help="search only the first N ready people")
    parser.add_argument("--check", action="store_true", help="read the sheet only, no CCIS")
    args = parser.parse_args()

    if not args.spreadsheet.is_file():
        print("File not found.")
        return 2
    if args.limit is not None and args.limit < 1:
        print("--limit must be 1 or more.")
        return 2

    try:
        if args.check:
            describe_layout(args.spreadsheet)
            describe_people(read_people(args.spreadsheet))
            return 0
        return run(args.spreadsheet, args.limit)
    except SystemProblem as error:
        print(f"Stopped: {error}")
        return 1


if __name__ == "__main__":
    sys.exit(main())
